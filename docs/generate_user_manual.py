from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "soft-copyright-manual"
IMG_DIR = OUT_DIR / "placeholders"
OUT_DOCX = OUT_DIR / "机器视觉智能问答与采集辅助系统用户使用手册.docx"

SYSTEM_NAME = "机器视觉智能问答与采集辅助系统"
MANUAL_TITLE = "用户使用手册"
SCHOOL = "武汉理工大学"
DATE_TEXT = "2026 年 7 月"


@dataclass(frozen=True)
class Figure:
    no: str
    caption: str
    replacement: str


FIGURES: list[Figure] = [
    Figure("图1.5.1-1", "后端服务启动成功", "后端终端启动成功截图，需显示端口 8123、context-path=/api 或 Spring Boot 启动完成信息"),
    Figure("图1.5.2-1", "前端服务启动成功", "前端 npm run dev 或生产构建成功截图，需显示访问地址"),
    Figure("图2.1-1", "用户注册界面", "注册页截图，需显示系统名称、注册页签、用户名/密码输入框和注册按钮"),
    Figure("图2.2-1", "用户登录界面", "登录页截图，需显示系统名称、登录页签、用户名/密码输入框和登录按钮"),
    Figure("图3.1.1-1", "系统首页", "登录后的首页截图，需显示三个入口：对话助手、采集工作台、知识库管理"),
    Figure("图3.1.2-1", "调试页面", "点击首页右上角设置按钮后的调试页截图，需显示后端地址、连接测试按钮和网络日志区域"),
    Figure("图3.2.1-1", "对话助手页面", "对话助手工作台截图，需显示左侧切换栏和视觉知识问答/视觉采集智能体区域"),
    Figure("图3.2.1-2", "视觉知识问答结果", "视觉知识问答截图，建议提问“什么是目标检测的 mAP”，需显示流式回复结果"),
    Figure("图3.2.2-1", "视觉采集智能体页面", "视觉采集智能体页面截图，需显示输入框和欢迎语"),
    Figure("图3.2.2-2", "智能体规划结果", "智能体根据采集或实验需求生成建议的结果截图"),
    Figure("图3.3.1-1", "知识库管理页面", "知识库管理页截图，需显示文档标题、文档内容输入框、上传按钮和文档列表"),
    Figure("图3.3.1-2", "知识文档上传结果", "上传一段机器视觉文档后的成功提示和已入库文档列表截图"),
    Figure("图3.4.1-1", "采集工作台页面", "采集工作台截图，需显示采集辅助规划、采集项目工作区、报告中心三个页签"),
    Figure("图3.4.1-2", "采集任务表单", "采集辅助规划表单截图，需显示任务类型、采集目标、采集地点、定位、时间预算、出行方式、归属项目"),
    Figure("图3.4.2-1", "采集规划结果页面", "点击生成采集规划后的结果截图，需显示计划标题、导出 PDF 按钮和点位列表"),
    Figure("图3.4.2-2", "多方案对比与点位详情", "采集规划结果中的方案 A/B/C、多选对比、点位得分、推荐理由、采集建议和风险提示截图"),
    Figure("图3.4.2-3", "地图与参考图信息", "点位卡片中高德静态地图、参考图或“在高德打开”导航链接截图"),
    Figure("图3.4.3-1", "采集计划 PDF 导出结果", "点击导出 PDF 后的成功提示截图，需显示生成报告路径或报告 ID"),
    Figure("图3.5.1-1", "新建采集项目", "采集项目工作区中新建项目区域截图，需显示项目名称、项目描述和创建按钮"),
    Figure("图3.5.2-1", "项目详情", "点击项目列表后的详情截图，需显示采集计划、实验计划、报告数量及下载链接"),
    Figure("图3.5.3-1", "阶段总结报告生成", "项目详情页点击生成报告后的成功提示截图"),
    Figure("图3.6.1-1", "报告中心列表", "报告中心表格截图，需显示标题、类型、所属项目、创建时间、下载 PDF"),
    Figure("图3.6.2-1", "PDF 报告下载结果", "浏览器下载 PDF 或打开 PDF 报告的截图，需显示文件已生成且中文正常"),
    Figure("图3.7.1-1", "实验流程规划对话", "在视觉采集智能体中输入实验规划需求后的对话结果截图"),
    Figure("图3.7.1-2", "实验计划 PDF 报告", "实验计划报告下载或打开截图，需显示数据准备、标注设计、训练步骤、评价指标等章节"),
]


def set_east_asian_font(run, font_name: str = "SimSun") -> None:
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")


def set_paragraph_font(paragraph, font_name: str = "SimSun", size: float = 12, bold: bool | None = None) -> None:
    for run in paragraph.runs:
        set_east_asian_font(run, font_name)
        run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold


def paragraph_border_bottom(paragraph, color: str = "000000", size: str = "8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def setup_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.line_spacing = 1.6
    normal.paragraph_format.space_after = Pt(4)

    for name, size, bold in [
        ("Heading 1", 16, True),
        ("Heading 2", 14, True),
        ("Heading 3", 13, True),
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
        style.paragraph_format.first_line_indent = None
        style.paragraph_format.space_before = Pt(13)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.5

    for name in ["Header", "Footer"]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(10.5)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")


def setup_section(section) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(3.17)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)


def setup_header_footer(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(SYSTEM_NAME)
    set_east_asian_font(run)
    run.font.size = Pt(10.5)
    paragraph_border_bottom(p, size="10")

    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_east_asian_font(run)
    run.font.size = Pt(10)


def add_center_text(doc: Document, text: str, size: float, bold: bool = False, after: float = 0) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_east_asian_font(run, "SimSun")
    run.font.size = Pt(size)
    run.bold = bold


def add_blank_lines(doc: Document, count: int, line_size: float = 12) -> None:
    for _ in range(count):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.space_after = Pt(0)
        p.add_run("").font.size = Pt(line_size)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.first_line_indent = None
    for run in p.runs:
        set_east_asian_font(run, "SimHei")
    return p


def add_para(doc: Document, text: str, first_line: bool = True, after: float = 4) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24) if first_line else None
    p.paragraph_format.line_spacing = 1.6
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_east_asian_font(run, "SimSun")
    run.font.size = Pt(12)
    return p


def add_item(doc: Document, text: str) -> None:
    p = add_para(doc, text, first_line=False, after=4)
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = None
    return p


def add_action(doc: Document, label: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"【{label}】")
    set_east_asian_font(run, "SimHei")
    run.font.size = Pt(12)
    run.bold = True


def draw_placeholder(fig: Figure, width=1050, height=590) -> Path:
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    path = IMG_DIR / f"{fig.no}.png"
    img = Image.new("RGB", (width, height), "#F4F5F7")
    draw = ImageDraw.Draw(img)
    try:
        cjk_font = "/System/Library/Fonts/STHeiti Medium.ttc"
        font_big = ImageFont.truetype(cjk_font, 42)
        font_mid = ImageFont.truetype(cjk_font, 30)
        font_small = ImageFont.truetype(cjk_font, 24)
    except Exception:
        font_big = ImageFont.load_default()
        font_mid = ImageFont.load_default()
        font_small = ImageFont.load_default()

    draw.rectangle((8, 8, width - 8, height - 8), outline="#7B8794", width=4)
    draw.rectangle((28, 28, width - 28, 90), fill="#E5E7EB", outline="#CBD5E1")
    draw.text((48, 44), SYSTEM_NAME, fill="#111827", font=font_mid)
    draw.text((width - 250, 44), fig.no, fill="#374151", font=font_mid)
    title = f"截图占位：{fig.caption}"
    tw = draw.textbbox((0, 0), title, font=font_big)[2]
    draw.text(((width - tw) / 2, 190), title, fill="#111827", font=font_big)
    replacement = "替换为：" + fig.replacement
    lines = wrap_text(draw, replacement, font_small, max_width=width - 130)
    y = 285
    for line in lines[:6]:
        lw = draw.textbbox((0, 0), line, font=font_small)[2]
        draw.text(((width - lw) / 2, y), line, fill="#4B5563", font=font_small)
        y += 42
    draw.rectangle((70, height - 92, width - 70, height - 44), outline="#CBD5E1", width=2)
    note = "请在 Word 中右键本图 -> 更改图片 -> 选择对应真实截图"
    nw = draw.textbbox((0, 0), note, font=font_small)[2]
    draw.text(((width - nw) / 2, height - 82), note, fill="#6B7280", font=font_small)
    img.save(path)
    return path


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    for ch in text:
        test = current + ch
        if draw.textbbox((0, 0), test, font=font)[2] <= max_width:
            current = test
        else:
            if current:
                lines.append(current)
            current = ch
    if current:
        lines.append(current)
    return lines


def add_figure(doc: Document, fig: Figure) -> None:
    path = draw_placeholder(fig)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run()
    r.add_picture(str(path), width=Inches(5.8))
    cap = doc.add_paragraph(f"{fig.no} {fig.caption}")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = None
    cap.paragraph_format.space_after = Pt(8)
    for run in cap.runs:
        set_east_asian_font(run, "SimSun")
        run.font.size = Pt(10.5)


def add_toc_line(doc: Document, title: str, page: int, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.left_indent = Pt((level - 1) * 24)
    p.paragraph_format.space_after = Pt(5 if level == 1 else 3)
    run = p.add_run(f"{title}{'.' * max(4, 74 - len(title) * 2 - level * 4)}{page}")
    set_east_asian_font(run, "SimSun")
    run.font.size = Pt(12 if level == 1 else 11)
    run.bold = level == 1
    if level == 3:
        run.italic = True


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "F2F2F2")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = None
        run = p.add_run(h)
        set_east_asian_font(run, "SimHei")
        run.bold = True
        run.font.size = Pt(10.5)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.first_line_indent = None
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i != 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            set_east_asian_font(run, "SimSun")
            run.font.size = Pt(10.5)
    doc.add_paragraph()


def add_replacement_table(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "4 截图替换清单", 1)
    add_para(doc, "本章列出手册中所有截图占位图对应的真实截图要求。正式提交软著材料前，建议按下表逐项替换为系统实际运行截图。")
    headers = ["图号", "图片位置", "需要替换的真实截图"]
    rows = [[fig.no, fig.caption, fig.replacement] for fig in FIGURES]
    add_table(doc, headers, rows, [1300, 2300, 5760])


def add_cover(doc: Document) -> None:
    add_blank_lines(doc, 9)
    add_center_text(doc, SYSTEM_NAME, 22, False)
    add_blank_lines(doc, 8)
    add_center_text(doc, MANUAL_TITLE, 24, False)
    add_blank_lines(doc, 5)
    add_center_text(doc, SCHOOL, 16, True)
    doc.add_page_break()
    add_blank_lines(doc, 2)
    add_center_text(doc, DATE_TEXT, 18, True)
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    add_center_text(doc, "目录", 14, False, after=18)
    toc_items = [
        ("1 系统简介与安装说明", 3, 1),
        ("1.1 引言", 3, 2),
        ("1.2 编写目的", 3, 2),
        ("1.3 系统概述", 3, 2),
        ("1.4 运行环境要求", 4, 2),
        ("1.5 系统安装说明", 4, 2),
        ("2 用户登录方法", 8, 1),
        ("2.1 用户注册", 8, 2),
        ("2.2 用户登录", 9, 2),
        ("2.3 退出系统", 10, 2),
        ("3 功能模块使用说明", 10, 1),
        ("3.1 首页与系统调试", 10, 2),
        ("3.2 对话助手", 12, 2),
        ("3.3 知识库管理", 15, 2),
        ("3.4 采集工作台", 17, 2),
        ("3.5 采集项目工作区", 22, 2),
        ("3.6 报告中心", 25, 2),
        ("3.7 视觉实验流程规划", 27, 2),
        ("3.8 系统数据与接口说明", 27, 2),
        ("4 截图替换清单", 28, 1),
    ]
    for title, page, level in toc_items:
        add_toc_line(doc, title, page, level)
    doc.add_page_break()


def add_install_section(doc: Document) -> None:
    add_heading(doc, "1 系统简介与安装说明", 1)
    add_heading(doc, "1.1 引言", 2)
    add_para(doc, "随着深度学习、智能感知与多模态数据处理技术的发展，机器视觉已广泛应用于目标检测、图像分割、深度估计、三维重建、自动驾驶感知、工业检测与科研实验等场景。机器视觉任务通常涉及知识查询、数据采集、场景选择、标注规范、实验流程设计、评价指标解释和阶段报告整理等多个环节，单纯依靠人工检索和手工记录容易出现效率低、流程不完整、材料沉淀不规范等问题。")
    add_para(doc, "本系统围绕机器视觉学习与实验场景，结合大语言模型、RAG 知识检索、智能体工具调用、地图与图片检索增强、PDF 报告生成和本地项目工作区能力，为用户提供从视觉知识问答到采集规划、实验规划和材料沉淀的一体化辅助服务。")
    add_heading(doc, "1.2 编写目的", 2)
    add_para(doc, "编写本用户使用手册的目的是为了全面阐述机器视觉智能问答与采集辅助系统的运行环境、安装部署方法、用户登录方式、功能模块和典型操作流程，使使用者能够快速掌握系统各项功能，并为系统演示、维护和后续升级提供必要说明。")
    add_heading(doc, "1.3 系统概述", 2)
    add_para(doc, "机器视觉智能问答与采集辅助系统采用前后端分离架构。后端基于 Spring Boot 3.4.4、Spring AI、Spring AI Alibaba 与 DashScope 大模型能力构建，提供知识问答、智能体对话、采集规划、实验规划、项目管理、报告生成和知识库管理接口；前端基于 Vue 3、Vue Router、Axios 与 SSE 实时通信构建，提供登录、首页、对话助手、采集工作台、知识库管理、项目工作区、报告中心和调试页面。")
    add_para(doc, "系统主要功能包括：用户注册登录、机器视觉知识问答、RAG 知识库增强问答、视觉采集智能体、场景化数据采集规划、多方案点位推荐、采集清单与隐私提示、视觉实验流程规划、采集项目归档、PDF 报告生成与下载、知识文档上传、后端连接诊断等。")
    add_heading(doc, "1.4 运行环境要求", 2)
    add_para(doc, "硬件环境要求：推荐使用 4 核及以上 CPU、8GB 及以上内存、20GB 及以上可用磁盘空间；如启用较大规模向量库、MCP 工具或本地文件处理，建议提高内存与磁盘配置。")
    add_para(doc, "软件环境要求：服务器端需安装 Java 17、Maven 或使用项目自带 mvnw、Spring Boot 运行环境；前端需安装 Node.js 与 npm；浏览器建议使用 Chrome、Edge、Firefox 或 Safari 的较新版本。")
    add_para(doc, "可选依赖包括 DashScope API Key、SEARCH_API_KEY、高德地图 API Key、Pexels API Key。DashScope 用于大模型问答；高德地图用于真实 POI、路线、静态地图和逆地理编码；Pexels 图片搜索用于补充采集点位参考图。未配置部分可选依赖时，系统会对部分能力进行降级，例如采集点位距离与多样性转为规则估算。")
    add_heading(doc, "1.5 系统安装说明", 2)
    add_heading(doc, "1.5.1 后端部署", 3)
    add_para(doc, "1. 进入项目根目录 cs-ai-agent，确认当前 Java 版本为 Java 17。若本机默认 Java 版本不是 17，需先设置 JAVA_HOME。")
    add_item(doc, "macOS 示例：export JAVA_HOME=$(/usr/libexec/java_home -v 17)")
    add_para(doc, "2. 配置必要环境变量。DASHSCOPE_API_KEY 为大模型调用所需；SEARCH_API_KEY、AMAP_MAPS_API_KEY、PEXELS_API_KEY 可按实际功能需要配置。")
    add_para(doc, "3. 执行后端编译命令，确认代码能够正常构建。")
    add_item(doc, "./mvnw -q compile")
    add_para(doc, "4. 启动 Spring Boot 后端服务。默认端口为 8123，context-path 为 /api。")
    add_item(doc, "./mvnw spring-boot:run")
    add_para(doc, "5. 启动成功后，可访问 http://localhost:8123/api/health 进行健康检查，返回 OK 表示后端服务可用，如图1.5.1-1所示。")
    add_figure(doc, FIGURES[0])
    add_heading(doc, "1.5.2 前端部署", 3)
    add_para(doc, "1. 进入前端目录 cs-ai-agent-fronted。")
    add_para(doc, "2. 安装前端依赖。")
    add_item(doc, "npm install")
    add_para(doc, "3. 本地开发时执行 npm run dev 启动 Vite 服务。前端默认通过 VITE_API_BASE_URL 或内置默认值 http://localhost:8123/api 访问后端接口。")
    add_item(doc, "npm run dev")
    add_para(doc, "4. 生产部署时可执行 npm run build 生成 dist 目录，并将构建产物部署到 Nginx 等静态资源服务器。若后端地址不是默认地址，应在环境变量或 .env 文件中配置 VITE_API_BASE_URL。")
    add_para(doc, "前端服务启动成功后，浏览器访问对应地址即可进入系统登录页面，如图1.5.2-1所示。")
    add_figure(doc, FIGURES[1])
    add_heading(doc, "1.5.3 MCP 与可选服务说明", 3)
    add_para(doc, "系统已在 mcp-servers.json 中配置高德地图 MCP 与图片搜索 MCP。高德地图 MCP 用于 POI 检索、路线规划、地图增强；图片搜索 MCP 用于补充采集点位参考图。首次启用图片搜索 MCP 前，需要先进入 cs-image-search-mcp-server 子模块构建 jar 包。")
    add_item(doc, "(cd cs-image-search-mcp-server && ../mvnw -q package -DskipTests)")
    add_para(doc, "如果仅需验证问答、采集规划规则兜底和报告等基础能力，可临时关闭 MCP Client。关闭后，采集规划仍可根据任务类型和场景规则生成候选点位，并在距离、可达性或多样性字段中标注估算。")
    add_item(doc, "./mvnw spring-boot:run -Dspring-boot.run.arguments='--spring.ai.mcp.client.enabled=false'")


def add_login_section(doc: Document) -> None:
    add_heading(doc, "2 用户登录方法", 1)
    add_heading(doc, "2.1 用户注册", 2)
    add_action(doc, "注册")
    add_para(doc, "用户首次使用系统时，输入系统访问地址进入登录页面，点击【注册】页签，填写用户名和密码后点击【注册】按钮。系统会校验用户名和密码是否为空，并检查用户名是否已存在。注册成功后页面提示“注册成功，请登录”，用户可返回登录页签继续登录，如图2.1-1所示。")
    add_figure(doc, FIGURES[2])
    add_heading(doc, "2.2 用户登录", 2)
    add_action(doc, "登录")
    add_para(doc, "用户在登录页签输入已注册的用户名和密码，点击【登录】按钮。系统验证通过后会返回 token 和用户名，前端将登录信息保存到本地并跳转至系统首页；若用户名或密码错误，页面会提示“用户名或密码错误”，用户需重新输入，如图2.2-1所示。")
    add_figure(doc, FIGURES[3])
    add_heading(doc, "2.3 退出系统", 2)
    add_action(doc, "退出")
    add_para(doc, "登录成功后，系统首页右上角显示当前用户名和【退出】按钮。用户点击【退出】后，前端会清除本地 token 和用户名，并跳转回登录页面。若用户未登录直接访问业务页面，系统路由会自动跳转至登录页面。")


def add_function_sections(doc: Document) -> None:
    add_heading(doc, "3 功能模块使用说明", 1)
    add_heading(doc, "3.1 首页与系统调试", 2)
    add_heading(doc, "3.1.1 首页功能入口", 3)
    add_para(doc, "（一）功能：系统首页用于展示机器视觉智能问答与采集辅助系统的主要业务入口，包括对话助手、采集工作台和知识库管理。用户可以在首页查看当前登录用户，也可以通过右上角设置按钮进入调试页面。")
    add_para(doc, "（二）使用方法：用户登录成功后默认进入首页。点击【对话助手】进入知识问答与智能体对话；点击【采集工作台】进入采集规划、项目工作区与报告中心；点击【知识库管理】进入文档上传与已入库文档查看页面，如图3.1.1-1所示。")
    add_figure(doc, FIGURES[4])
    add_heading(doc, "3.1.2 调试页面", 3)
    add_para(doc, "（一）功能：调试页面用于查看当前后端接口地址、默认端口、前端地址配置、MCP 本地开关说明，并提供后端连接测试、视觉问答 SSE 测试、视觉智能体 SSE 测试、网络请求日志和手动消息测试能力。")
    add_para(doc, "（二）使用方法：点击首页右上角设置按钮进入调试页面。用户可点击【测试基本连接】确认 /health 接口是否可用；点击【测试视觉问答SSE】验证知识问答流式接口；点击【测试视觉智能体SSE】验证智能体流式接口；也可在手动测试区域输入测试消息并查看响应内容，如图3.1.2-1所示。")
    add_figure(doc, FIGURES[5])

    add_heading(doc, "3.2 对话助手", 2)
    add_heading(doc, "3.2.1 视觉知识问答", 3)
    add_para(doc, "（一）功能：视觉知识问答模块面向机器视觉概念、目标检测、图像分割、深度估计、三维重建、数据采集、标注规范、实验流程和评价指标等问题提供问答服务。系统通过 SSE 流式响应逐步显示回答，并为每个会话自动生成会话 ID，支持多轮上下文记忆。")
    add_para(doc, "（二）使用方法：点击首页【对话助手】，在左侧选择【视觉知识问答】，在输入框中输入机器视觉相关问题后点击【发送】。系统会在对话区域实时输出回答，AI 回复支持 Markdown 渲染并经过安全净化处理，如图3.2.1-1和图3.2.1-2所示。")
    add_figure(doc, FIGURES[6])
    add_action(doc, "发送")
    add_para(doc, "用户可输入“目标检测和图像分割有什么区别”“mAP@0.5 和 mAP@0.5:0.95 怎么算”“小目标检测有哪些难点”等问题。系统会结合内置机器视觉知识库和大模型能力输出解释。")
    add_figure(doc, FIGURES[7])
    add_heading(doc, "3.2.2 视觉采集智能体", 3)
    add_para(doc, "（一）功能：视觉采集智能体用于处理复杂机器视觉任务，可通过对话方式理解用户的采集或实验需求，并根据任务调用采集规划、实验规划、搜索、网页抓取、资源下载、文件操作、PDF 生成等工具，生成可执行建议。")
    add_para(doc, "（二）使用方法：在对话助手左侧选择【视觉采集智能体】，输入采集或实验需求，例如“在上海静安区做车辆和行人目标检测采集，半天，步行”或“帮我生成目标检测实验流程规划”。系统会在内部调用合适工具并输出最终结果，如图3.2.2-1和图3.2.2-2所示。")
    add_figure(doc, FIGURES[8])
    add_action(doc, "智能规划")
    add_para(doc, "智能体输出时不展示内部工具名称、工具参数和原始返回，仅向用户展示整理后的采集建议、实验步骤、报告路径或必要提示。若请求与机器视觉无关，系统会说明功能边界并进行简要引导。")
    add_figure(doc, FIGURES[9])

    add_heading(doc, "3.3 知识库管理", 2)
    add_heading(doc, "3.3.1 上传知识文档", 3)
    add_para(doc, "（一）功能：知识库管理模块用于向系统知识库补充机器视觉领域文档。用户可输入文档标题和正文内容，系统将文本加入向量库，并在后续 RAG 问答中作为知识来源。")
    add_para(doc, "（二）使用方法：点击首页【知识库管理】，在【文档标题】输入框填写文档名称，在【文档内容】输入框粘贴机器视觉资料、实验规范或项目说明，点击【上传到知识库】按钮。上传成功后页面提示“上传成功！文档已入库”，如图3.3.1-1和图3.3.1-2所示。")
    add_figure(doc, FIGURES[10])
    add_figure(doc, FIGURES[11])
    add_heading(doc, "3.3.2 文档列表与知识问答增强", 3)
    add_para(doc, "（一）功能：知识库页面右侧展示已入库文档列表，包括标题、领域和正文预览。系统启动时会加载 resources/document 下的机器视觉文档，例如机器视觉基础知识、目标检测常见问题、图像分割指南、深度估计与三维重建指南、无人驾驶视觉感知、评价指标说明和采集安全隐私规范。")
    add_para(doc, "（二）使用方法：点击【刷新】按钮可重新获取当前内存向量库中的文档列表。用户上传的文档属于 manual 领域，适合补充课程资料、实验说明和项目要求。当前版本知识库主要采用内存向量库，服务重启后会重新加载内置资源文档，生产环境可按需启用 PgVector 等持久化向量存储方案。")

    add_heading(doc, "3.4 采集工作台", 2)
    add_heading(doc, "3.4.1 采集辅助规划", 3)
    add_para(doc, "（一）功能：采集辅助规划是系统核心功能之一。用户填写任务类型、采集目标、采集地点、时间预算、出行方式和归属项目后，系统会生成采集计划，包括推荐点位、多套采集方案、路线与时段建议、采集清单、标注指引、隐私与安全提示。")
    add_para(doc, "（二）使用方法：点击首页【采集工作台】，默认进入【采集辅助规划】页签。用户可选择任务类型（目标检测、图像分割、深度估计、三维重建、图像分类、目标跟踪），输入采集目标（如车辆、行人、交通灯），填写采集地点或点击定位按钮自动解析当前位置，选择时间预算和出行方式，并可选择归属项目，如图3.4.1-1和图3.4.1-2所示。")
    add_figure(doc, FIGURES[12])
    add_action(doc, "生成采集规划")
    add_para(doc, "点击【生成采集规划】按钮后，前端向 /vision/collection/plan 提交结构化任务。后端会先规范化视觉任务，再优先调用高德地图 MCP 检索真实 POI；若地图能力不可用，则使用规则候选点兜底，保证接口稳定返回。")
    add_figure(doc, FIGURES[13])
    add_heading(doc, "3.4.2 采集方案对比与点位查看", 3)
    add_para(doc, "（一）功能：系统会对候选采集点位进行场景分类和四维打分，评分维度包括场景匹配、可达性、多样性和安全性，总分 0-100。推荐结果支持多方案展示，包括质量优先、多样性优先和效率优先。每个点位包含名称、地址、距离、场景标签、适合任务、得分、推荐理由、采集建议和风险提示。")
    add_para(doc, "（二）使用方法：采集规划生成后，页面右侧显示计划结果。用户可点击方案页签选择或多选方案进行对比，查看每个点位的得分、行程距离、预计耗时、场景标签和安全提示。若点位具有经纬度，系统会补充高德静态地图、导航链接和参考图，如图3.4.2-1、图3.4.2-2和图3.4.2-3所示。")
    add_figure(doc, FIGURES[14])
    add_figure(doc, FIGURES[15])
    add_figure(doc, FIGURES[16])
    add_heading(doc, "3.4.3 采集规划 PDF 导出", 3)
    add_para(doc, "（一）功能：系统可将采集规划导出为 PDF 报告，报告内容包括任务概述、推荐采集点位、路线与时段建议、采集清单、标注与隐私提示。PDF 文件统一生成到 tmp/vision/reports 目录，并在报告中心生成记录。")
    add_para(doc, "（二）使用方法：在采集规划结果区域点击【导出 PDF】按钮。系统会根据当前采集计划 ID 调用 /vision/report/pdf 接口，生成 collection_plan 类型报告。生成成功后页面显示报告路径或报告 ID，如图3.4.3-1所示。")
    add_figure(doc, FIGURES[17])

    add_heading(doc, "3.5 采集项目工作区", 2)
    add_heading(doc, "3.5.1 新建采集项目", 3)
    add_para(doc, "（一）功能：采集项目工作区用于按项目组织采集计划、实验计划和报告。用户可创建项目，后续在采集规划表单中选择归属项目，实现采集结果和报告的统一归档。")
    add_para(doc, "（二）使用方法：在采集工作台左侧选择【采集项目工作区】，填写项目名称和项目描述后点击【创建项目】按钮。创建成功后项目会出现在项目列表中，如图3.5.1-1所示。")
    add_figure(doc, FIGURES[18])
    add_heading(doc, "3.5.2 查看项目详情与关联记录", 3)
    add_para(doc, "（一）功能：项目详情页聚合展示项目元数据、关联采集计划、关联实验计划和报告记录。用户可通过项目列表查看每个项目下的采集、实验和报告数量。")
    add_para(doc, "（二）使用方法：点击项目列表中的某个项目，系统调用 /vision/project/{id} 接口获取详情。详情区域展示采集计划标题、实验计划任务类型、报告标题和 PDF 下载链接，如图3.5.2-1所示。")
    add_figure(doc, FIGURES[19])
    add_heading(doc, "3.5.3 生成阶段总结报告", 3)
    add_para(doc, "（一）功能：项目详情页提供阶段总结报告生成能力。用户点击【生成报告】后，系统会根据项目名称和描述生成 stage_summary 类型 PDF 报告，并写入报告中心。")
    add_para(doc, "（二）使用方法：进入项目详情后点击【生成报告】按钮。生成成功后页面提示“已生成报告，可在报告中心查看”，如图3.5.3-1所示。")
    add_figure(doc, FIGURES[20])

    add_heading(doc, "3.6 报告中心", 2)
    add_heading(doc, "3.6.1 查看报告列表", 3)
    add_para(doc, "（一）功能：报告中心用于统一查看系统已生成的 PDF 报告记录。报告类型包括采集计划、实验计划、调研摘要和阶段总结。列表展示报告标题、报告类型、所属项目、创建时间和操作入口。")
    add_para(doc, "（二）使用方法：在采集工作台左侧选择【报告中心】，系统会同时加载报告列表和项目列表，将项目 ID 映射为项目名称。点击【刷新】可重新获取最新报告记录，如图3.6.1-1所示。")
    add_figure(doc, FIGURES[21])
    add_heading(doc, "3.6.2 下载 PDF 报告", 3)
    add_para(doc, "（一）功能：报告中心支持下载每条 PDF 报告。系统根据报告 ID 定位本地 PDF 文件，设置 application/pdf 响应类型和下载文件名后返回文件。")
    add_para(doc, "（二）使用方法：在报告列表中点击某条记录右侧【下载 PDF】链接，浏览器会打开或下载对应 PDF 文件。若文件不存在或报告 ID 不存在，系统返回 404。如图3.6.2-1所示。")
    add_figure(doc, FIGURES[22])

    add_heading(doc, "3.7 视觉实验流程规划", 2)
    add_heading(doc, "3.7.1 实验流程生成", 3)
    add_para(doc, "（一）功能：实验流程规划模块根据任务类型和实验目标生成可执行实验计划。当前系统内置目标检测、图像分割、深度估计、三维重建四类模板，输出内容包括数据准备、标注设计、数据集划分、推荐模型、训练步骤、评价指标和风险注意事项。")
    add_para(doc, "（二）使用方法：用户可在【对话助手】-【视觉采集智能体】中输入实验规划需求，例如“帮我生成目标检测实验流程，目标是完成城市道路车辆和行人检测实验”。智能体会调用实验规划工具生成计划，如图3.7.1-1所示。")
    add_figure(doc, FIGURES[23])
    add_action(doc, "导出实验计划")
    add_para(doc, "实验计划生成并落库后，可通过报告接口生成 experiment_plan 类型 PDF 报告。报告章节包括实验目标、数据准备、标注设计、数据集划分、推荐模型、训练步骤、评价指标、风险与注意事项，如图3.7.1-2所示。")
    add_figure(doc, FIGURES[24])

    add_heading(doc, "3.8 系统数据与接口说明", 2)
    add_heading(doc, "3.8.1 本地数据存储", 3)
    add_para(doc, "系统采用轻量本地 JSON 持久化保存采集计划、实验计划、项目和报告索引。默认存储根目录为 tmp/vision，其中 plans.json 保存采集计划，experiments.json 保存实验计划，projects.json 保存采集项目，records.json 保存报告索引，reports 目录保存 PDF 报告文件。")
    add_heading(doc, "3.8.2 业务接口响应格式", 3)
    add_para(doc, "系统接口采用双轨响应格式。/api/ai/** 聊天与智能体接口返回纯文本或 SSE 流式响应，便于前端逐 token 展示；/api/vision/** 业务接口统一返回 ApiResponse<T>，字段包括 code、message 和 data。其中 code=0 表示成功，40001 表示参数校验失败，40400 表示资源不存在，50000 表示服务内部错误，50001 表示依赖能力不可用。")
    add_heading(doc, "3.8.3 依赖能力与降级策略", 3)
    add_para(doc, "当 DashScope API Key 未配置时，大模型问答和智能体能力不可用；当高德地图或 MCP 不可用时，采集规划会自动降级为规则候选点和估算评分，仍可返回采集点位、清单和隐私提示；当图片搜索不可用时，点位参考图为空或仅展示高德实拍图，不影响采集规划正文信息。")


def build_doc() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_styles(doc)
    setup_section(doc.sections[0])
    setup_header_footer(doc.sections[0])
    add_cover(doc)
    add_toc(doc)
    add_install_section(doc)
    add_login_section(doc)
    add_function_sections(doc)
    add_replacement_table(doc)
    doc.save(OUT_DOCX)


if __name__ == "__main__":
    build_doc()
    print(OUT_DOCX)
